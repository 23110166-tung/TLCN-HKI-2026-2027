# -*- coding: utf-8 -*-
"""
Script to generate a professional Progress Report Word Document (.docx) and PDF (.pdf)
for the Graduation / Specialization Thesis (Tiểu luận chuyên ngành AI).
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from pipeline.pdf_converter import PDFConverter

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, top="CCCCCC", bottom="CCCCCC", left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = ['<w:tcBorders ' + nsdecls('w') + '>']
    if top:
        borders.append(f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>')
    else:
        borders.append('<w:top w:val="none"/>')
    if bottom:
        borders.append(f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{bottom}"/>')
    else:
        borders.append('<w:bottom w:val="none"/>')
    if left:
        borders.append(f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>')
    else:
        borders.append('<w:left w:val="none"/>')
    if right:
        borders.append(f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>')
    else:
        borders.append('<w:right w:val="none"/>')
    borders.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(borders)))

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102) # Dark Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87)
    return p

def add_body_p(doc, text, bold=False, italic=False, bullet=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=3.5, line_spacing=1.15)
    if bullet:
        p.paragraph_format.left_indent = Inches(0.25)
        run_b = p.add_run("• ")
        run_b.font.name = 'Times New Roman'
        run_b.font.bold = True
        run_b.font.color.rgb = RGBColor(0, 51, 102)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.italic = italic
    return p

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    out_docx = os.path.join(BASE_DIR, "Bao_Cao_Tien_Do_Tieu_Luan_Chuyen_Nganh.docx")
    out_pdf = os.path.join(BASE_DIR, "Bao_Cao_Tien_Do_Tieu_Luan_Chuyen_Nganh.pdf")
    
    doc = docx.Document()
    
    # Page Margins (Standard A4: Top/Bottom 2cm, Left 2.5cm, Right 2cm)
    for sec in doc.sections:
        sec.top_margin = Inches(0.79)
        sec.bottom_margin = Inches(0.79)
        sec.left_margin = Inches(0.98)
        sec.right_margin = Inches(0.79)
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        
    # Header Info Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    # Left Cell
    cell_l = header_table.rows[0].cells[0]
    cell_l.width = Inches(3.8)
    p_l = cell_l.paragraphs[0]
    format_paragraph(p_l, space_before=0, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r1 = p_l.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(10)
    r2 = p_l.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP.HCM\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    r2.font.bold = True
    r3 = p_l.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(10)
    r3.font.bold = True
    r4 = p_l.add_run("---o0o---")
    r4.font.name = 'Times New Roman'
    r4.font.size = Pt(9)
    
    # Right Cell
    cell_r = header_table.rows[0].cells[1]
    cell_r.width = Inches(3.8)
    p_r = cell_r.paragraphs[0]
    format_paragraph(p_r, space_before=0, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_c1 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r_c1.font.name = 'Times New Roman'
    r_c1.font.size = Pt(10)
    r_c1.font.bold = True
    r_c2 = p_r.add_run("Độc lập – Tự do – Hạnh phúc\n")
    r_c2.font.name = 'Times New Roman'
    r_c2.font.size = Pt(10)
    r_c2.font.bold = True
    r_c3 = p_r.add_run("--------------------")
    r_c3.font.name = 'Times New Roman'
    r_c3.font.size = Pt(9)
    
    # Title
    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=14, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_t = p_title.add_run("BÁO CÁO TIẾN ĐỘ THỰC HIỆN TIỂU LUẬN CHUYÊN NGÀNH AI")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 51, 102)
    
    p_sub = doc.add_paragraph()
    format_paragraph(p_sub, space_before=0, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_s = p_sub.add_run("ĐỢT 1: BÁO CÁO ĐỊNH KỲ 2 TUẦN (GIAI ĐOẠN THU THẬP, GÁN NHÃN & THIẾT LẬP DỮ LIỆU)")
    r_s.font.name = 'Times New Roman'
    r_s.font.size = Pt(11)
    r_s.font.italic = True
    r_s.font.bold = True
    
    # General Info Box
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    rows_data = [
        ("Tên đề tài:", "XÂY DỰNG HỆ THỐNG NHẬN DẠNG, TRÍCH XUẤT VÀ KIỂM TRA TÍNH HỢP LỆ CỦA THÔNG TIN TỪ ĐƠN HỌC VỤ VÀ BẢNG ĐIỂM SINH VIÊN BẰNG HỌC SÂU"),
        ("Mã môn học:", "GRPR471979 – Tiểu luận chuyên ngành AI"),
        ("Giảng viên hướng dẫn:", "TS. Bùi Mạnh Quân"),
        ("Nhóm sinh viên thực hiện:", "1. Nguyễn Bách Tùng – MSSV: 23110166\n2. Nguyễn Hưng Nguyên – MSSV: 23110035\n3. Trần Cẩm Long – MSSV: 23110122"),
        ("Thời gian thực hiện:", "Học kỳ 1 – Năm học 2026-2027 (Đợt báo cáo tiến độ: Tuần 1 & 2)")
    ]
    
    for idx, (label, val) in enumerate(rows_data):
        row = info_table.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(5.4)
        set_cell_background(c1, "F2F5F9" if idx % 2 == 0 else "FFFFFF")
        set_cell_background(c2, "F2F5F9" if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        set_cell_border(c1, top="D0D7DE", bottom="D0D7DE", left="D0D7DE", right=None)
        set_cell_border(c2, top="D0D7DE", bottom="D0D7DE", left=None, right="D0D7DE")
        
        p1 = c1.paragraphs[0]
        format_paragraph(p1, space_before=0, space_after=0)
        r_lbl = p1.add_run(label)
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.bold = True
        
        p2 = c2.paragraphs[0]
        format_paragraph(p2, space_before=0, space_after=0)
        r_val = p2.add_run(val)
        r_val.font.name = 'Times New Roman'
        r_val.font.size = Pt(10.5)
        if idx == 0:
            r_val.font.bold = True
            r_val.font.color.rgb = RGBColor(0, 51, 102)

    # Divider line
    p_div = doc.add_paragraph()
    format_paragraph(p_div, space_before=8, space_after=8)
    
    # -------------------------------------------------------------
    # SECTION I
    # -------------------------------------------------------------
    add_heading_1(doc, "I. TỔNG KẾT CÁC CÔNG VIỆC ĐÃ HOÀN THÀNH (MILESTONE 1)")
    add_body_p(doc, "Bám sát các chỉ đạo và yêu cầu của Thầy hướng dẫn về việc phân công công việc, thu thập biểu mẫu chính thức từ nhà trường và xây dựng bộ dữ liệu gán nhãn có bảo mật thông tin sinh viên theo Đề cương chi tiết, nhóm đã hoàn thành các hạng mục công việc sau:")
    
    add_heading_2(doc, "1. Thu thập & Chuẩn hóa Tài liệu, Biểu mẫu Học vụ HCMUTE")
    add_body_p(doc, "Nguồn thu thập: Nhóm đã tra cứu, liên hệ và tổng hợp các biểu mẫu quản lý chất lượng ISO hiện hành từ Cổng thông tin Phòng Đào tạo (pdt.hcmute.edu.vn), các Khoa chuyên môn và hệ thống quản lý đào tạo trực tuyến của nhà trường.", bullet=True)
    add_body_p(doc, "Đã chuẩn hóa 10 loại tài liệu trọng tâm phục vụ đề tài:", bullet=True)
    
    # 10 Templates table
    t_docs = doc.add_table(rows=11, cols=4)
    t_docs.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_docs.autofit = False
    
    headers = ["STT", "Mã Biểu Mẫu", "Tên Loại Tài Liệu / Biểu Mẫu", "Đặc Điểm Nghiệp Vụ Xử Lý"]
    col_widths = [Inches(0.5), Inches(1.3), Inches(3.2), Inches(2.6)]
    
    # Format Table Header
    hdr_row = t_docs.rows[0]
    for c_i, name in enumerate(headers):
        cell = hdr_row.cells[c_i]
        cell.width = col_widths[c_i]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        set_cell_border(cell, top="003366", bottom="003366", left="003366", right="003366")
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(name)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    doc_rows = [
        ("1", "T01_DIEM_I", "Đơn xin nhận và trả điểm I (Hoãn thi)", "Bao gồm 5 chữ ký viết tay: SV, GV 1, GV 2, Khoa, PĐT."),
        ("2", "T02_HOC_2CTDT", "Đơn đăng ký học cùng lúc hai CTĐT", "Sinh viên từ năm 2 đăng ký song bằng, GPA >= 2.5."),
        ("3", "T03_CHUYEN_HE", "Đơn xin chuyển hệ đào tạo", "Chuyển hệ CQ - VLVH (Mộc BGH + Mộc PĐT)."),
        ("4", "T04_CHUYEN_DIEM", "Đơn xin chuyển điểm học phần", "Chuyển đổi điểm các học phần tương đương."),
        ("5", "T05_BAO_LUU", "Đơn xin bảo lưu điểm môn học", "Bảo lưu kết quả học tập khi chuyển chương trình."),
        ("6", "T06_BAN_SAO", "Đơn xin cấp bản sao từ sổ gốc", "Cựu sinh viên đã tốt nghiệp xin cấp bản sao văn bằng."),
        ("7", "T07_CHUYEN_NN", "Đơn chuyển điểm Ngoại ngữ tốt nghiệp", "Nộp chứng chỉ TOEIC/IELTS xét chuẩn đầu ra tốt nghiệp."),
        ("8", "T08_THAY_THE", "Đơn xin học môn thay thế tốt nghiệp", "Mộc Khoa đóng chồng lên chữ ký Khoa + Mộc PĐT."),
        ("9", "T09_MIEN_THI", "Đơn xin miễn thi học phần", "Miễn thi các học phần đã tích lũy ở trường cũ."),
        ("10", "T10_TRANSCRIPT", "Bảng điểm học tập trực tuyến Portal", "Cấu trúc bảng điểm đa học kỳ của 6 ngành đào tạo.")
    ]
    
    for r_i, r_data in enumerate(doc_rows):
        row = t_docs.rows[r_i + 1]
        bg_color = "F7F9FC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(r_data):
            cell = row.cells[c_i]
            cell.width = col_widths[c_i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            set_cell_border(cell, top="E1E4E8", bottom="E1E4E8", left="E1E4E8", right="E1E4E8")
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 1] else WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(p, space_before=0, space_after=0, align=align)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            if c_i == 1:
                r.font.bold = True
                
    # Subsection 2
    add_heading_2(doc, "2. Xây dựng Pipeline Sinh Dữ Liệu Tự Động & Gán Nhãn Chuẩn YOLO")
    add_body_p(doc, "Thiết kế Không gian Nhãn (Ontology): Xây dựng bộ ontology gồm 18 lớp học vụ (student_name, student_id, class_name, faculty, major, course_name, course_code, credits, score, stamp, signature, date, decision_number...).", bullet=True)
    add_body_p(doc, "Logic phân bổ Niên khóa chính xác (K20 đến K26): Đảm bảo tính nhất quán nghiệp vụ: K26 là tân sinh viên Năm 1 (vừa nhập học HK01 2026-2027), K25 (Năm 2), K24 (Năm 3), K23 (Năm 4) và K22-K20 (sinh viên tốt nghiệp).", bullet=True)
    add_body_p(doc, "Cơ chế Bảo mật và Ẩn danh hóa (Theo Mục 4 Đề cương): Dữ liệu được sinh bằng thuật toán Synthetic Data tự động hóa 100%, không sử dụng thông tin nhạy cảm thật của sinh viên (số CCCD, tài khoản ngân hàng, địa chỉ riêng); loại bỏ toàn bộ metadata EXIF của ảnh.", bullet=True)
    add_body_p(doc, "Quy mô Bộ Dữ liệu đã gán nhãn: Hoàn thành 500 ảnh chất lượng cao (450 ảnh biểu mẫu + 50 ảnh bảng điểm trực tuyến) với hơn 22.500 bounding box chuẩn xác pixel, chia tỉ lệ chuẩn 80% Train – 10% Validation – 10% Test.", bullet=True)

    # Subsection 3
    add_heading_2(doc, "3. Đóng Gói & Công Bố Tập Dữ Liệu Lên Nền Tảng Kaggle")
    add_body_p(doc, "Toàn bộ 500 ảnh và các file nhãn YOLO text đã được nhóm đóng gói chuẩn hóa và tải lên nền tảng Kaggle để sẵn sàng kết nối huấn luyện trên môi trường GPU Cloud:")
    add_body_p(doc, "Liên kết Kaggle Dataset: https://www.kaggle.com/datasets/nguyenbachtung/database-hcmute", bold=True, italic=True)

    # Subsection 4
    add_heading_2(doc, "4. Chuẩn Bị Bộ Dữ Liệu In Ấn Benchmark Thực Tế (100 Mẫu A4)")
    add_body_p(doc, "Xuất bản 100 File PDF A4 300 DPI: Nhóm đã tạo thư mục print_dataset_100/ chứa 100 file PDF riêng lẻ và 1 file Master gộp ALL_100_PAGES_MASTER.pdf.", bullet=True)
    add_body_p(doc, "Nhúng sẵn Con dấu mộc tròn đỏ & Chữ ký tay: 100% tài liệu được nhúng sẵn con dấu đỏ chính quy (Phòng Đào tạo, Ban Giám hiệu, Khoa) và chữ ký mực xanh sinh động đúng vị trí thực tế.", bullet=True)
    add_body_p(doc, "Ground Truth Manifest (dataset_manifest.json): Lưu trữ toàn bộ siêu dữ liệu (Họ tên, MSSV, Khóa, Ngành, Điểm số, Môn học) của cả 100 mẫu để tự động đối soát.", bullet=True)

    # Subsection 5
    add_heading_2(doc, "5. Hoàn Thành Thu Thập, Xử Lý Nắn Thẳng & Gán Nhãn 120 Ảnh Chụp Thực Tế")
    add_body_p(doc, "Thu thập 120 ảnh chụp smartphone thực tế: Đã in ấn các biểu mẫu A4 và chụp lại bằng camera điện thoại trên mặt bàn gỗ với đầy đủ các điều kiện nhiễu thực tế (góc nghiêng 10°-25°, bóng đổ, ánh sáng phòng, độ cong mép giấy).", bullet=True)
    add_body_p(doc, "Xây dựng thuật toán Nắn thẳng phối cảnh (Perspective Dewarping): Tự động phát hiện 4 điểm cực trị của viền giấy A4 và nắn phẳng về chuẩn trực diện 2480 x 3508 (300 DPI).", bullet=True)
    add_body_p(doc, "Căn chỉnh tương quan chéo & Gán nhãn tự động 100%: Tự động đo lường độ lệch in ấn (dx, dy) và gán nhãn thành công 9.172 Bounding Box chuẩn YOLOv11 cho toàn bộ 120 ảnh thực tế trong dataset_yolo_real_world/.", bullet=True)
    add_body_p(doc, "Trích xuất 360+ mẫu ảnh cắt vùng chữ (Crops OCR): Chuẩn bị sẵn sàng các crop ảnh chữ in, chữ viết tay, mã số, điểm số phục vụ huấn luyện mô hình VietOCR.", bullet=True)

    # -------------------------------------------------------------
    # SECTION II
    # -------------------------------------------------------------
    add_heading_1(doc, "II. BẢNG PHÂN CÔNG NHIỆM VỤ CỤ THỂ TRONG NHÓM")
    add_body_p(doc, "Theo định hướng phân công công việc rõ ràng giữa các thành viên, nhóm đã phân chia trách nhiệm cụ thể như sau:")
    
    t_assign = doc.add_table(rows=4, cols=4)
    t_assign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_assign.autofit = False
    
    hdr_assign = ["STT", "Thành Viên Thực Hiện", "Nhiệm Vụ Đảm Nhận", "Tiến Độ / Kết Quả"]
    w_assign = [Inches(0.5), Inches(1.8), Inches(3.3), Inches(2.0)]
    
    for c_i, name in enumerate(hdr_assign):
        cell = t_assign.rows[0].cells[c_i]
        cell.width = w_assign[c_i]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        set_cell_border(cell, top="003366", bottom="003366", left="003366", right="003366")
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(name)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    assign_rows = [
        ("1", "Nguyễn Bách Tùng\n(MSSV: 23110166)", 
         "• Thiết kế Ontology 18 lớp & Xây dựng Pipeline sinh dữ liệu tự động.\n• Thu thập và chuẩn hóa dữ liệu CTĐT 6 ngành.\n• Đóng gói và phát hành Dataset lên Kaggle.\n• Cài đặt môi trường huấn luyện YOLOv11.", 
         "Đã hoàn thành 100%\n(Dataset 500 mẫu + Pipeline tự động)"),
        
        ("2", "Nguyễn Hưng Nguyên\n(MSSV: 23110035)", 
         "• Nghiên cứu tài liệu các mô hình Layout Detection (YOLOv8, YOLOv11) và Vision-Language.\n• Đã in ấn bộ biểu mẫu A4 và tổ chức chụp 120 ảnh thực tế bằng camera smartphone.\n• Xây dựng pipeline nắn thẳng Dewarping và hoàn tất gán nhãn 120 ảnh thực tế.", 
         "Đã hoàn thành 100%\n(120 ảnh chụp thực tế +\nGán nhãn 9.172 Box YOLO)"),
        
        ("3", "Trần Cẩm Long\n(MSSV: 23110122)", 
         "• Nghiên cứu kiến trúc nhận dạng chữ tiếng Việt (VietOCR, TrOCR).\n• Khảo sát xây dựng Rule Engine kiểm tra tính hợp lệ (Regex, Logic MSSV, Khóa học).\n• Chuẩn bị giao diện Web Demo (FastAPI / Streamlit).", 
         "Đã hoàn thành khảo sát lý thuyết &\nkhung kiến trúc Rule Engine")
    ]
    
    for r_i, r_data in enumerate(assign_rows):
        row = t_assign.rows[r_i + 1]
        bg_color = "F7F9FC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(r_data):
            cell = row.cells[c_i]
            cell.width = w_assign[c_i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            set_cell_border(cell, top="E1E4E8", bottom="E1E4E8", left="E1E4E8", right="E1E4E8")
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(p, space_before=0, space_after=0, align=align)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            if c_i == 1:
                r.font.bold = True

    # -------------------------------------------------------------
    # SECTION III
    # -------------------------------------------------------------
    add_heading_1(doc, "III. KẾ HOẠCH THỰC HIỆN TRONG 2 TUẦN TỚI (MILESTONE 2)")
    add_body_p(doc, "1. Huấn luyện Mô hình Cơ sở (Baseline Training) trên Kaggle GPU:", bold=True)
    add_body_p(doc, "Tiến hành huấn luyện mô hình YOLOv11 (nano / small) trên tập 500 ảnh synthetic đã chuẩn hóa 18 lớp ontology.", bullet=True)
    add_body_p(doc, "Đo lường và theo dõi các chỉ số định lượng: Precision, Recall, mAP@0.5, mAP@0.5:0.95 trên tập Validation và Test.", bullet=True)
    
    add_body_p(doc, "2. Đánh giá Zero-Shot & Phân tích Hiện tượng Lệch Miền (Domain Shift):", bold=True)
    add_body_p(doc, "Đưa trực tiếp 120 ảnh chụp camera thực tế đã gán nhãn vào mô hình đã train để đánh giá độ chính xác thực tế và đo lường độ suy giảm hiệu năng do nhiễu môi trường.", bullet=True)
    add_body_p(doc, "Áp dụng các kỹ thuật Data Augmentation nâng cao (Photometric Distortions, Random Perspective, Motion Blur) để thu hẹp khoảng cách miền.", bullet=True)
    
    add_body_p(doc, "3. Triển khai Mô-đun Nhận dạng Chữ Tiếng Việt (VietOCR / TrOCR) & Ghép Cặp Key-Value:", bold=True)
    add_body_p(doc, "Thực nghiệm nhận dạng ký tự tiếng Việt trên 360+ ảnh crop text line của tập thực tế.", bullet=True)
    add_body_p(doc, "Xây dựng thuật toán ghép nối trường thông tin theo cấu trúc không gian (Spatial Key-Value Association).", bullet=True)

    add_body_p(doc, "4. Xây dựng Rule Engine Kiểm Tra Tính Hợp Lệ & Nhất Quán Học Vụ:", bold=True)
    add_body_p(doc, "Thiết kế bộ quy tắc tự động: Kiểm tra định dạng MSSV theo niên khóa, đối soát điều kiện điểm trung bình, xác thực sự tồn tại của con dấu tròn đỏ và chữ ký tay hợp lệ.", bullet=True)
    
    add_body_p(doc, "5. Thiết Kế & Hoàn Thiện Giao Diện Web Demo (FastAPI / Streamlit):", bold=True)
    add_body_p(doc, "Xây dựng ứng dụng web cho phép người dùng tải ảnh biểu mẫu học vụ, hiển thị trực quan các Bounding Box nhận dạng, bảng trích xuất dữ liệu và kết quả kiểm tra tính hợp lệ.", bullet=True)

    # -------------------------------------------------------------
    # SIGNATURE SECTION
    # -------------------------------------------------------------
    p_sig_space = doc.add_paragraph()
    format_paragraph(p_sig_space, space_before=10, space_after=6)
    
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    cell_s1 = sig_table.rows[0].cells[0]
    cell_s1.width = Inches(3.8)
    p_s1 = cell_s1.paragraphs[0]
    format_paragraph(p_s1, space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_g = p_s1.add_run("Ý KIẾN CỦA GIẢNG VIÊN HƯỚNG DẪN\n")
    r_g.font.name = 'Times New Roman'
    r_g.font.size = Pt(11)
    r_g.font.bold = True
    r_g2 = p_s1.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n\n")
    r_g2.font.name = 'Times New Roman'
    r_g2.font.size = Pt(10)
    r_g2.font.italic = True
    r_g3 = p_s1.add_run("TS. Bùi Mạnh Quân")
    r_g3.font.name = 'Times New Roman'
    r_g3.font.size = Pt(11)
    r_g3.font.bold = True
    
    cell_s2 = sig_table.rows[0].cells[1]
    cell_s2.width = Inches(3.8)
    p_s2 = cell_s2.paragraphs[0]
    format_paragraph(p_s2, space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_sv1 = p_s2.add_run("Tp. Hồ Chí Minh, ngày 31 tháng 08 năm 2026\n")
    r_sv1.font.name = 'Times New Roman'
    r_sv1.font.size = Pt(10)
    r_sv1.font.italic = True
    r_sv2 = p_s2.add_run("ĐẠI DIỆN NHÓM SINH VIÊN\n")
    r_sv2.font.name = 'Times New Roman'
    r_sv2.font.size = Pt(11)
    r_sv2.font.bold = True
    r_sv3 = p_s2.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n\n")
    r_sv3.font.name = 'Times New Roman'
    r_sv3.font.size = Pt(10)
    r_sv3.font.italic = True
    r_sv4 = p_s2.add_run("Nguyễn Bách Tùng\nNguyễn Hưng Nguyên – Trần Cẩm Long")
    r_sv4.font.name = 'Times New Roman'
    r_sv4.font.size = Pt(11)
    r_sv4.font.bold = True

    # Save docx with lock fallback
    actual_docx = out_docx
    actual_pdf = out_pdf
    try:
        doc.save(out_docx)
    except PermissionError:
        actual_docx = os.path.join(BASE_DIR, "Bao_Cao_Tien_Do_Tieu_Luan_Chuyen_Nganh_Cap_Nhat.docx")
        actual_pdf = os.path.join(BASE_DIR, "Bao_Cao_Tien_Do_Tieu_Luan_Chuyen_Nganh_Cap_Nhat.pdf")
        doc.save(actual_docx)
        
    print(f"✅ Đã tạo file Word báo cáo tiến độ: {actual_docx}")
    
    # Convert to PDF
    try:
        conv = PDFConverter()
        conv.convert(actual_docx, actual_pdf)
        conv.close()
        print(f"✅ Đã xuất file PDF báo cáo tiến độ: {actual_pdf}")
    except Exception as e:
        print(f"⚠️ Lỗi convert PDF: {e}")

if __name__ == "__main__":
    main()
