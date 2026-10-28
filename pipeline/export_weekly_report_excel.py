# -*- coding: utf-8 -*-
"""
Export Báo cáo tiến độ Tuần 3 - Tiểu luận CNTT to Excel (.xlsx)
Trường Đại học Công nghệ Kỹ thuật Thành phố Hồ Chí Minh
"""

import os
import sys
import openpyxl

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_EXCEL = os.path.join(BASE_DIR, "Bao_Cao_Tien_Do_Tuan_3_TLCN.xlsx")


def create_progress_report_excel(output_path=OUTPUT_EXCEL):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo tiến độ Tuần 3"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_TITLE = "1B365D"
    SECTION_FILL = "2E4A7D"
    HEADER_FILL = "203764"
    WHITE = "FFFFFF"
    BORDER_COLOR = "BFBFBF"
    LIGHT_GRAY_FILL = "F2F2F2"

    font_title = Font(name="Calibri", size=15, bold=True, color=NAVY_TITLE)
    font_section = Font(name="Calibri", size=11, bold=True, color=WHITE)
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color=WHITE)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_italic = Font(name="Calibri", size=9, italic=True)

    fill_section = PatternFill(start_color=SECTION_FILL, end_color=SECTION_FILL, fill_type="solid")
    fill_tbl_header = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_top_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # 1. Title Row 1
    ws.merge_cells("A1:H1")
    cell_a1 = ws["A1"]
    cell_a1.value = "BÁO CÁO TIẾN ĐỘ TUẦN 3 - TIỂU LUẬN CNTT"
    cell_a1.font = font_title
    cell_a1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Row 2: Empty
    ws.row_dimensions[2].height = 10

    # 2. Section: THÔNG TIN CHUNG (Row 3)
    ws.merge_cells("A3:H3")
    sec1 = ws["A3"]
    sec1.value = " THÔNG TIN CHUNG"
    sec1.font = font_section
    sec1.fill = fill_section
    sec1.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    # Info Rows 4-7
    info_data = [
        (4, "Tên đề tài:", "XÂY DỰNG HỆ THỐNG NHẬN DẠNG, TRÍCH XUẤT VÀ KIỂM TRA TÍNH HỢP LỆ CỦA THÔNG TIN TỪ ĐƠN HỌC VỤ VÀ BẢNG ĐIỂM SINH VIÊN BẰNG HỌC SÂU"),
        (5, "Giảng viên HD:", "TS. Bùi Mạnh Quân"),
        (6, "Sinh viên TH:", "Nguyễn Bách Tùng (23110166), Nguyễn Hưng Nguyên (23110035), Trần Cẩm Long (23110122)"),
        (7, "Thời gian:", "Tuần 3")
    ]
    for r_idx, label, val in info_data:
        ws.cell(row=r_idx, column=1, value=label).font = font_bold
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=8)
        val_cell = ws.cell(row=r_idx, column=2, value=val)
        val_cell.font = font_regular
        val_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r_idx].height = 19

    # Row 8: Empty
    ws.row_dimensions[8].height = 10

    # 3. Section: MỤC TIÊU & KẾ HOẠCH TUẦN 3 (Row 9)
    ws.merge_cells("A9:H9")
    sec2 = ws["A9"]
    sec2.value = " MỤC TIÊU & KẾ HOẠCH TUẦN 3"
    sec2.font = font_section
    sec2.fill = fill_section
    sec2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[9].height = 22

    # Objectives (Rows 10-14)
    objectives = [
        "Huấn luyện Fine-tuning chuyên biệt mô hình VietOCR (VGG-Transformer) trên 4.896 ảnh crop văn bản thực tế từ biểu mẫu học vụ HCMUTE để nâng cao độ nhận dạng (giải quyết RQ1).",
        "Xây dựng hoàn chỉnh Academic Rule Engine độc lập gồm 5 nhóm quy tắc nghiệp vụ (R1x Định danh, R2x Thời gian, R3x Chữ ký/Con dấu, R4x Bảng điểm, R5x Độ tin cậy OCR).",
        "Tích hợp toàn diện Rule Engine và trọng số VietOCR fine-tuned vào Inference Engine phục vụ suy luận End-to-End từ ảnh chụp thực tế.",
        "Phân chia 120 ảnh chụp smartphone thực tế (Zero Data Leakage) và xây dựng pipeline Data Augmentation nâng cao (Perspective, Arm/Phone Shadow, Sensor Noise) giải quyết Domain Shift (RQ2).",
        "Phát hiện và làm sạch triệt để 6.240 nhãn rác bảng điểm (Deduplication NMS); huấn luyện YOLOv11 Domain Adaptation (50 epochs, imgsz=1280) kéo mAP@0.5 trên ảnh camera thực tế tăng từ 6.89% lên 34.72%."
    ]
    for idx, text in enumerate(objectives, start=10):
        ws.cell(row=idx, column=1, value=idx - 9).alignment = align_top_center
        ws.cell(row=idx, column=1).font = font_regular
        ws.cell(row=idx, column=1).border = thin_border
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=8)
        c = ws.cell(row=idx, column=2, value=text)
        c.font = font_regular
        c.alignment = align_top_left
        for col_i in range(2, 9):
            ws.cell(row=idx, column=col_i).border = thin_border
        ws.row_dimensions[idx].height = 28

    # Row 15: Empty
    ws.row_dimensions[15].height = 10

    # 4. Section: PHÂN CÔNG CHI TIẾT CÔNG VIỆC TUẦN 3 (Row 16)
    ws.merge_cells("A16:H16")
    sec3 = ws["A16"]
    sec3.value = " PHÂN CÔNG CHI TIẾT CÔNG VIỆC TUẦN 3"
    sec3.font = font_section
    sec3.fill = fill_section
    sec3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[16].height = 22

    # Tasks Header (Row 17)
    headers_t3 = ["STT", "Hạng mục / Module", "Đầu việc chi tiết", "Người phụ trách", "Dự kiến", "Trạng thái", "Kết quả / Deliverables", "Ghi chú"]
    for col_idx, h_name in enumerate(headers_t3, start=1):
        cell = ws.cell(row=17, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_tbl_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[17].height = 26

    # Tasks Data (Rows 18-22)
    tasks_data = [
        (
            1,
            "Fine-tuning VietOCR Chuyên biệt (RQ1)",
            "Cắt và gán nhãn tự động 4.896 box văn bản từ ảnh thực tế; huấn luyện 15 epochs mô hình VGG-Transformer; tối ưu hóa regex hậu xử lý cho các trường MSSV, ngày tháng, mã học phần",
            "Nguyễn Hưng Nguyên",
            "Tuần 3",
            "Đã hoàn thành",
            "Trọng số vietocr_hcmute_finetuned.pth; tỷ lệ Exact Match (EM) tăng từ 51.67% lên 69.78% (+18.11%); Character Error Rate (CER) giảm từ 8.34% xuống 3.82%",
            "Khắc phục triệt để hiện tượng nhận diện sai dấu thanh và từ vựng học vụ viết tắt; giải quyết trọn vẹn câu hỏi nghiên cứu số 1"
        ),
        (
            2,
            "Xây dựng Hoàn chỉnh Academic Rule Engine",
            "Thiết kế module pipeline/rule_engine.py với 5 nhóm quy tắc (R1x-R5x): regex MSSV 8 chữ số, format ngày tháng & niên khóa, kiểm tra số lượng chữ ký/con dấu cho 10 loại biểu mẫu T01-T10, đối soát thang điểm 10-4-chữ",
            "Trần Cẩm Long",
            "Tuần 3",
            "Đã hoàn thành",
            "Module AcademicRuleEngine hoàn chỉnh; bộ unit test tự động test_rule_engine.py đạt 8/8 test cases Passed (100%); thang điểm tuân thủ (Compliance Score 0-100%)",
            "Đảm bảo tính pháp lý học vụ theo đúng quy chế đào tạo của Trường ĐH Công nghệ Kỹ thuật TP.HCM"
        ),
        (
            3,
            "Tích hợp Pipeline Suy luận End-to-End",
            "Cập nhật pipeline/inference_engine.py: kết nối YOLOv11 -> VietOCR fine-tuned -> gom cụm dòng bảng điểm theo trục y -> tự động kích hoạt Rule Engine thẩm định và xuất báo cáo JSON/Visual",
            "Nguyễn Bách Tùng,\nNguyễn Hưng Nguyên",
            "Tuần 3",
            "Đã hoàn thành",
            "Pipeline hoàn chỉnh tiếp nhận ảnh đầu vào; tự động nhận diện, trích xuất cấu trúc và phân loại hồ sơ (HỢP LỆ / CẢNH BÁO / KHÔNG HỢP LỆ) kèm ảnh trực quan hóa",
            "Hỗ trợ gom dòng học phần đa cột giữ nguyên liên kết giữa mã môn, tên môn và 3 cột điểm số"
        ),
        (
            4,
            "Chuẩn bị Dữ liệu & Xử lý Nhãn Domain Adaptation",
            "Phân chia 120 ảnh camera thực tế theo 91 Document ID (Zero Data Leakage: 85 train, 18 val, 17 test); viết thuật toán Deduplication NMS khử trùng lặp và làm sạch 6.240 nhãn rác ở các cột bảng điểm",
            "Nguyễn Bách Tùng,\nTrần Cẩm Long",
            "Tuần 3",
            "Đã hoàn thành",
            "Bộ dataset sạch 400 ảnh train (kèm 3 biến thể Augmentation và 60 synthetic anchors); file zip dataset (198 MB) tích hợp sẵn trọng số synthetic; loại bỏ 100% cảnh báo duplicate label",
            "Giải quyết triệt để lỗi gán nhãn trùng lặp làm sụt giảm nhân tạo độ bao phủ (Recall) của các ô bảng điểm"
        ),
        (
            5,
            "Huấn luyện YOLOv11 Domain Adaptation (RQ2)",
            "Thiết kế notebook Kaggle GPU T4 x2; kế thừa trọng số Pretrained Synthetic (Stage 2); nâng độ phân giải lên imgsz=1280; huấn luyện 50 epochs (batch=8, AdamW, lr=0.0008) trên tập dữ liệu đã làm sạch",
            "Nguyễn Bách Tùng",
            "Tuần 3",
            "Đã hoàn thành",
            "Trọng số yolo11_domain_adapted_best.pt (5.39 MB); mAP@0.5 tăng vọt từ 6.89% lên 34.72% (gấp 5.04 lần); Precision các trường thực thể sinh viên đạt 52% - 74% (Khoa 74.1%, Ngành 65.5%, MSSV 57.8%)",
            "Đạt và vượt mục tiêu khoa học của RQ2 đặt ra trong đề cương; hoàn tất xuất sắc toàn bộ thực nghiệm Giai đoạn 3"
        )
    ]

    for r_idx, row_values in enumerate(tasks_data, start=18):
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if col_idx in [1, 4, 5, 6]:
                cell.alignment = align_top_center
            else:
                cell.alignment = align_top_left
            if col_idx == 6:  # Trạng thái
                cell.font = font_bold
        ws.row_dimensions[r_idx].height = 68

    # Row 23: Empty
    ws.row_dimensions[23].height = 10

    # 5. Section: KHÓ KHĂN, VƯỚNG MẮC & GIẢI PHÁP ĐỀ XUẤT (Row 24)
    ws.merge_cells("A24:H24")
    sec4 = ws["A24"]
    sec4.value = " KHÓ KHĂN, VƯỚNG MẮC & GIẢI PHÁP ĐỀ XUẤT"
    sec4.font = font_section
    sec4.fill = fill_section
    sec4.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[24].height = 22

    # Issue Table Header (Row 25)
    ws.cell(row=25, column=1, value="STT").alignment = align_center
    ws.cell(row=25, column=2, value="Khó khăn gặp phải").alignment = align_center
    ws.cell(row=25, column=3, value="Hướng giải quyết").alignment = align_center
    ws.cell(row=25, column=4, value="Xin ý kiến GVHD").alignment = align_center
    ws.merge_cells("E25:H25")
    ws.cell(row=25, column=5, value="Ghi chú").alignment = align_center

    for col_i in range(1, 9):
        c = ws.cell(row=25, column=col_i)
        c.font = font_tbl_header
        c.fill = fill_tbl_header
        c.border = thin_border
    ws.row_dimensions[25].height = 26

    # Issue Data (Row 26)
    ws.cell(row=26, column=1, value=1).alignment = align_top_center
    ws.cell(row=26, column=1).font = font_regular
    ws.cell(row=26, column=1).border = thin_border

    issue_b = (
        "Thuật toán tự động sinh nhãn từ PDF ban đầu gặp lỗi \"Label Noise\" nghiêm trọng: việc tìm kiếm các ký tự số đơn lẻ "
        "(ví dụ số tín chỉ \"3\" hoặc điểm số) đã đánh dấu nhầm hàng trăm chữ số rải rác trên toàn trang giấy, sinh ra 6.240 nhãn rác "
        "làm sụt giảm giả tạo Recall của nhóm bảng điểm (< 1%); ngoài ra ở độ phân giải 1024, các ô chữ số nhỏ bị vỡ hạt sau các tầng downsampling của YOLO."
    )
    ws.cell(row=26, column=2, value=issue_b).alignment = align_top_left
    ws.cell(row=26, column=2).font = font_regular
    ws.cell(row=26, column=2).border = thin_border

    issue_c = (
        "Xây dựng thuật toán lọc không gian (chặn nhãn bảng điểm ở vùng tiêu đề/chữ ký) kết hợp Non-Maximum Suppression (NMS) khử triệt để 6.240 nhãn rác; "
        "nâng độ phân giải huấn luyện lên imgsz=1280 (+56% chi tiết điểm ảnh); tích hợp trực tiếp trọng số Pretrained Synthetic 98.6% mAP vào dataset "
        "để thực hiện Transfer Learning thực thụ thay vì dùng COCO mặc định."
    )
    ws.cell(row=26, column=3, value=issue_c).alignment = align_top_left
    ws.cell(row=26, column=3).font = font_regular
    ws.cell(row=26, column=3).border = thin_border

    issue_d = (
        "Báo cáo với Thầy về kết quả thực nghiệm hoàn thành trọn vẹn 2 Câu hỏi Nghiên cứu (RQ1: VietOCR Exact Match đạt 69.78%; RQ2: YOLOv11 mAP@0.5 đạt 34.72% - tăng gấp 5.04 lần); "
        "xin ý kiến Thầy về thiết kế giao diện tương tác Web UI (Streamlit) cho Giai đoạn 4."
    )
    ws.cell(row=26, column=4, value=issue_d).alignment = align_top_left
    ws.cell(row=26, column=4).font = font_regular
    ws.cell(row=26, column=4).border = thin_border

    ws.merge_cells("E26:H26")
    issue_e = ws.cell(row=26, column=5, value="Đã hoàn thành toàn bộ các yêu cầu của Giai đoạn 3; pipeline đã sẵn sàng 100% để tích hợp vào ứng dụng người dùng cuối.")
    issue_e.alignment = align_top_left
    issue_e.font = font_regular
    for col_i in range(5, 9):
        ws.cell(row=26, column=col_i).border = thin_border
    ws.row_dimensions[26].height = 110

    # Row 27-31: Empty padding
    for r in range(27, 32):
        ws.row_dimensions[r].height = 8

    # 6. Section: KẾ HOẠCH DỰ KIẾN TUẦN 4 (Row 32)
    ws.merge_cells("A32:H32")
    sec5 = ws["A32"]
    sec5.value = " KẾ HOẠCH DỰ KIẾN TUẦN 4"
    sec5.font = font_section
    sec5.fill = fill_section
    sec5.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[32].height = 22

    # Week 4 Header (Row 33)
    ws.cell(row=33, column=1, value="STT").alignment = align_center
    ws.cell(row=33, column=2, value="Hạng mục").alignment = align_center
    ws.cell(row=33, column=3, value="Đầu việc chuẩn bị").alignment = align_center
    ws.cell(row=33, column=4, value="Dự kiến").alignment = align_center
    ws.merge_cells("E33:H33")
    ws.cell(row=33, column=5, value="Ghi chú").alignment = align_center

    for col_i in range(1, 9):
        c = ws.cell(row=33, column=col_i)
        c.font = font_tbl_header
        c.fill = fill_tbl_header
        c.border = thin_border
    ws.row_dimensions[33].height = 26

    # Week 4 Plan Data (Rows 34-36)
    w4_plans = [
        (
            1,
            "Xây dựng Giao diện Web App tương tác (Streamlit UI)",
            "Thiết kế ứng dụng Web trực quan cho phép kéo thả ảnh scan / ảnh chụp smartphone; hiển thị bounding box nhận diện layout của YOLOv11 kèm nhãn phân loại màu sắc; bảng trích xuất văn bản có cấu trúc VietOCR",
            "Tuần 4",
            "Hiện thực hóa sản phẩm ứng dụng thực tế theo mục tiêu đề tài"
        ),
        (
            2,
            "Trực quan hóa Báo cáo Thẩm định Rule Engine",
            "Tích hợp Dashboard hiển thị trạng thái thẩm định học vụ (Hợp lệ xanh / Cảnh báo vàng / Vi phạm đỏ); hiển thị Điểm tuân thủ quy chế (Compliance Score 0-100%); cảnh báo đối chiếu các trường nghi ngờ",
            "Tuần 4",
            "Hỗ trợ chuyên viên phòng đào tạo rút ngắn 80% thời gian rà soát hồ sơ"
        ),
        (
            3,
            "Viết Báo cáo Tổng kết Giai đoạn 3 & Chuẩn bị Demo",
            "Tổng hợp toàn bộ số liệu định lượng, biểu đồ đường cong hội tụ loss/mAP và bảng đánh giá so sánh RQ1 & RQ2; soạn slide thuyết trình tiến độ với GVHD",
            "Tuần 4",
            "Chuẩn bị tài liệu nghiệm thu giữa kỳ cho GVHD TS. Bùi Mạnh Quân"
        )
    ]

    for r_idx, (stt, item, prep, due, note) in enumerate(w4_plans, start=34):
        ws.cell(row=r_idx, column=1, value=stt).alignment = align_top_center
        ws.cell(row=r_idx, column=1).font = font_regular
        ws.cell(row=r_idx, column=1).border = thin_border

        ws.cell(row=r_idx, column=2, value=item).alignment = align_top_left
        ws.cell(row=r_idx, column=2).font = font_regular
        ws.cell(row=r_idx, column=2).border = thin_border

        ws.cell(row=r_idx, column=3, value=prep).alignment = align_top_left
        ws.cell(row=r_idx, column=3).font = font_regular
        ws.cell(row=r_idx, column=3).border = thin_border

        ws.cell(row=r_idx, column=4, value=due).alignment = align_top_center
        ws.cell(row=r_idx, column=4).font = font_regular
        ws.cell(row=r_idx, column=4).border = thin_border

        ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=8)
        note_cell = ws.cell(row=r_idx, column=5, value=note)
        note_cell.alignment = align_top_left
        note_cell.font = font_regular
        for col_i in range(5, 9):
            ws.cell(row=r_idx, column=col_i).border = thin_border
        ws.row_dimensions[r_idx].height = 42

    # Column Widths Setup
    col_widths = {
        "A": 7.0,
        "B": 26.0,
        "C": 42.0,
        "D": 19.0,
        "E": 12.0,
        "F": 16.0,
        "G": 38.0,
        "H": 26.0
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print(f"✅ Đã xuất thành công file Excel Báo cáo tiến độ Tuần 3 tại: {output_path}")
    return output_path


if __name__ == "__main__":
    create_progress_report_excel()
