# -*- coding: utf-8 -*-
"""
DOCX Template Filler for Vietnamese Academic Forms
Fills paragraphs and tables with realistic mock data and tracks inserted strings with class labels.
"""

import os
import re
import random
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def clean_dots(text):
    """Replaces sequences of dots, spaced dots, underscores with empty space."""
    return re.sub(r'(\.\s*){3,}|[…_]{3,}', '', text)

def fill_table_cell(cell, text, cls_name, recorded_list, font_size_pt=8.5):
    """Fills a table cell cleanly and records the field."""
    text_clean = str(text).strip()
    cell.text = text_clean
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(font_size_pt)
    if text_clean and cls_name:
        recorded_list.append({"text": text_clean, "class_name": cls_name})

def fill_document(template_path, student_data):
    """
    Fills a given docx template according to its specific structure and patterns.
    Returns: (doc, recorded_fields)
    """
    filename = os.path.basename(template_path)
    doc = docx.Document(template_path)
    recorded_fields = []
    st = student_data
    
    if "BAN SAO TU SO GOC" in filename:
        _fill_ban_sao(doc, st, recorded_fields)
    elif "DonXinChuyenDiemNgoaiNgu" in filename or "BM10" in filename:
        _fill_chuyen_diem_nn(doc, st, recorded_fields)
    elif "ThayThe" in filename or "BM12" in filename:
        _fill_mon_thay_the(doc, st, recorded_fields)
    elif "Don mien thi" in filename:
        _fill_mien_thi(doc, st, recorded_fields)
    elif "Don xin bao luu" in filename:
        _fill_bao_luu(doc, st, recorded_fields)
    elif "DonXinChuyenDiem_2018" in filename:
        _fill_chuyen_diem_2018(doc, st, recorded_fields)
    elif "Don_xin_chuyenhedaotao" in filename:
        _fill_chuyen_he(doc, st, recorded_fields)
    elif "hoc2CTDT" in filename or "TBaovaMaudangky" in filename:
        _fill_hoc_2_ctdt(doc, st, recorded_fields)
    elif "nhận và trả điểm I" in filename or "diem I" in filename:
        _fill_diem_i(doc, st, recorded_fields)
    else:
        _fill_generic(doc, st, recorded_fields)
        
    return doc, recorded_fields


def _fill_ban_sao(doc, st, rec):
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Tôi tên:" in t:
            p.text = f"            Tôi tên:    {st['student_name']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
        elif "Sinh ngày:" in t and "Tại:" in t:
            p.text = f"            Sinh ngày:      {st['dob']}        Tại:    {st['pob']}"
            rec.append({"text": st['dob'], "class_name": "dob"})
            rec.append({"text": st['pob'], "class_name": "pob"})
        elif "Ngành học:" in t and "Khóa học:" in t:
            p.text = f"            Ngành học:        {st['major']}         Khóa học:   2021 - 2025"
            rec.append({"text": st['major'], "class_name": "major"})
            rec.append({"text": "2021 - 2025", "class_name": "degree_info"})
        elif "Hệ đào tạo:" in t:
            p.text = f"            Hệ đào tạo:         Đại học chính quy"
            rec.append({"text": "Đại học chính quy", "class_name": "degree_info"})
        elif "Năm tốt nghiệp:" in t and "Ngày cấp bằng:" in t:
            p.text = f"            Năm tốt nghiệp: 2025      Ngày cấp bằng: {st['cert_date']}"
            rec.append({"text": "2025", "class_name": "date"})
            rec.append({"text": st['cert_date'], "class_name": "date"})
        elif "Số hiệu bằng:" in t and "Số vào sổ:" in t:
            p.text = f"Số hiệu bằng: {st['degree_number']}         Số vào sổ: {st['book_number']}"
            rec.append({"text": st['degree_number'], "class_name": "degree_info"})
            rec.append({"text": st['book_number'], "class_name": "degree_info"})
        elif "Nay tôi làm đơn này xin cấp Bản sao:" in t:
            p.text = f"Nay tôi làm đơn này xin cấp Bản sao: Bằng tốt nghiệp đại học"
            rec.append({"text": "Bằng tốt nghiệp đại học", "class_name": "degree_info"})
        elif "Số lượng cấp:" in t:
            p.text = f"Số lượng cấp: 03 bản"
            rec.append({"text": "03 bản", "class_name": "degree_info"})
        elif "Lý do:" in t:
            p.text = f"            Lý do: {st['reason']}"
            rec.append({"text": st['reason'], "class_name": "reason"})
        elif "Số đt liên hệ:" in t:
            p.text = f" Số đt liên hệ: {st['phone']}"
            rec.append({"text": st['phone'], "class_name": "phone"})
        elif "ngày" in t and "tháng" in t and "năm 20" in t:
            p.text = f"                                                                              Tp. Hồ Chí Minh, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)


def _fill_chuyen_diem_nn(doc, st, rec):
    # Set compact margins to guarantee single page fit
    for sec in doc.sections:
        sec.top_margin = Inches(0.4)
        sec.bottom_margin = Inches(0.35)
        sec.left_margin = Inches(0.6)
        sec.right_margin = Inches(0.6)

    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Họ tên sinh viên:" in t:
            p.text = f"Họ tên sinh viên:  {st['student_name']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
        elif "MSSV:" in t and "Điện thoại:" in t:
            p.text = f"MSSV:  {st['student_id']} \tĐiện thoại:  {st['phone']}"
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['phone'], "class_name": "phone"})
        elif "Lớp:" in t and "Khoa:" in t:
            p.text = f"Lớp:  {st['class_name']} \tKhoa:  {st['faculty']}"
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Đã thi đạt chứng chỉ:" in t:
            cert_title = st.get('cert_name', 'TOEIC Quốc tế (IIG Việt Nam)')
            p.text = f"Đã thi đạt chứng chỉ: {cert_title}"
            rec.append({"text": cert_title, "class_name": "degree_info"})
        elif "Với số điểm:" in t or "Nghe:" in t:
            p.text = f"Với số điểm: Nghe: {st['cert_listening']} Nói: {st['cert_speaking']} Đọc: {st['cert_reading']} Viết: {st['cert_writing']} Tổng điểm: {st['cert_total']}"
            if st['cert_listening'] != "---":
                rec.append({"text": st['cert_listening'], "class_name": "score"})
            if st['cert_speaking'] != "---":
                rec.append({"text": st['cert_speaking'], "class_name": "score"})
            if st['cert_reading'] != "---":
                rec.append({"text": st['cert_reading'], "class_name": "score"})
            if st['cert_writing'] != "---":
                rec.append({"text": st['cert_writing'], "class_name": "score"})
            rec.append({"text": str(st['cert_total']), "class_name": "score"})
        elif "Tại:" in t and "Ngày cấp chứng chỉ:" in t:
            p.text = f"Tại: {st['cert_place']} Ngày cấp chứng chỉ: {st['cert_date']}"
            rec.append({"text": st['cert_place'], "class_name": "pob"})
            rec.append({"text": st['cert_date'], "class_name": "date"})
        elif "Tp. HCM, ngày" in t:
            p.text = f"                                                                      Tp. HCM, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif "Xác nhận của Phòng Đào tạo" in t:
            p.text = "Xác nhận của Phòng Đào tạo                                         Sinh viên ký tên"
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)

    if doc.tables:
        t = doc.tables[0]
        cert_type = st.get("cert_type", "TOEIC")
        cert_total = st.get("cert_total", "600")

        # Map strictly according to Phu Luc I - Quy dinh chuyen diem tieng Anh DH SPKT TP.HCM
        if cert_type == "TOEIC":
            try:
                score = int(cert_total)
            except:
                score = 600
                
            if score >= 650:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "10")
                ]
            elif score >= 600:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "7.5")
                ]
            elif score >= 550:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "7.5"),
                    ("Anh văn 5", "ENGL530537", "3", "5.0")
                ]
            elif score >= 500:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "7.5"),
                    ("Anh văn 4", "ENGL430437", "3", "5.0")
                ]
            elif score >= 455:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "8.5"),
                    ("Anh văn 3", "ENGL330337", "3", "6.0")
                ]
            else:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "7.5"),
                    ("Anh văn 3", "ENGL330337", "3", "5.0")
                ]
        else: # IELTS
            try:
                score = float(cert_total)
            except:
                score = 6.0
                
            if score >= 6.0:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "10")
                ]
            elif score >= 5.5:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "8.5")
                ]
            elif score >= 5.0:
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "8.0")
                ]
            else: # IELTS 4.5
                eng_courses = [
                    ("Anh văn 1", "ENGL130137", "3", "10"),
                    ("Anh văn 2", "ENGL230237", "3", "10"),
                    ("Anh văn 3", "ENGL330337", "3", "10"),
                    ("Anh văn 4", "ENGL430437", "3", "10"),
                    ("Anh văn 5", "ENGL530537", "3", "8.0")
                ]

        for row_idx in range(1, len(t.rows)):
            row = t.rows[row_idx]
            c_idx = row_idx - 1
            
            # Set trHeight to 340 dxa (17pt)
            if row._tr.trPr is not None:
                for child in list(row._tr.trPr):
                    if child.tag.endswith('trHeight'):
                        row._tr.trPr.remove(child)
            row._tr.get_or_add_trPr().append(parse_xml(r'<w:trHeight %s w:val="340" w:hRule="atLeast"/>' % nsdecls('w')))

            if c_idx < len(eng_courses):
                c_name, c_code, c_cred, c_score = eng_courses[c_idx]
                vals = [str(row_idx), c_name, c_code, c_cred, c_score]
                classes = [None, "course_name", "course_code", "credits", "score"]
            else:
                vals = [str(row_idx), "", "", "", ""]
                classes = [None, None, None, None, None]

            for col_idx in range(5):
                cell = row.cells[col_idx]
                val = vals[col_idx]
                cls = classes[col_idx]
                
                p = cell.paragraphs[0]
                p.text = val
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if col_idx in [0, 2, 3, 4]:
                    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.5)
                
                if val and cls:
                    rec.append({"text": val, "class_name": cls})


def _fill_mon_thay_the(doc, st, rec):
    # Set compact margins to guarantee single page fit
    for sec in doc.sections:
        sec.top_margin = Inches(0.5)
        sec.bottom_margin = Inches(0.45)
        sec.left_margin = Inches(0.65)
        sec.right_margin = Inches(0.65)

    # 3 distinct courses for replacement logic
    c1 = st["courses"][0]
    c2 = st["courses"][1] if len(st["courses"]) > 1 else st["courses"][0]
    c3 = st["courses"][2] if len(st["courses"]) > 2 else None

    # Track how many course lines we filled
    course_line_idx = 0

    for i, p in enumerate(doc.paragraphs):
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text

        if "Khoa" in t and "Kính gửi" not in t and len(t) < 40:
            p.text = f"		- Khoa: {st['faculty']}"
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Em tên:" in t and "Mã số SV:" in t:
            p.text = f"Em tên:  {st['student_name']} \t\t Mã số SV:  {st['student_id']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['student_id'], "class_name": "student_id"})
        elif "Lớp:" in t and "Điện thoại:" in t:
            p.text = f"Lớp:  {st['class_name']} \t Điện thoại:  {st['phone']} \t Email:  {st['email']}"
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['phone'], "class_name": "phone"})
            rec.append({"text": st['email'], "class_name": "email"})
        elif "Tên môn:" in t and "Mã MH:" in t and "Số tín chỉ:" in t:
            if course_line_idx == 0:
                # Môn trong CTĐT cần thay thế
                p.text = f"Tên môn:  {c1['name']} \t Mã MH:  {c1['code']} \t Số tín chỉ:  {c1['credits']}"
                rec.append({"text": c1['name'], "class_name": "course_name"})
                rec.append({"text": c1['code'], "class_name": "course_code"})
                rec.append({"text": str(c1['credits']), "class_name": "credits"})
                course_line_idx += 1
            elif course_line_idx == 1:
                # Môn thay thế 1
                p.text = f"Tên môn:  {c2['name']} \t Mã MH:  {c2['code']} \t Số tín chỉ:  {c2['credits']}"
                rec.append({"text": c2['name'], "class_name": "course_name"})
                rec.append({"text": c2['code'], "class_name": "course_code"})
                rec.append({"text": str(c2['credits']), "class_name": "credits"})
                course_line_idx += 1
            elif course_line_idx == 2:
                # Môn thay thế 2
                if c3:
                    p.text = f"và Tên môn:  {c3['name']} \t Mã MH:  {c3['code']} \t Số tín chỉ:  {c3['credits']}"
                    rec.append({"text": c3['name'], "class_name": "course_name"})
                    rec.append({"text": c3['code'], "class_name": "course_code"})
                    rec.append({"text": str(c3['credits']), "class_name": "credits"})
                else:
                    p.text = "và Tên môn: \t Mã MH: \t Số tín chỉ:"
                course_line_idx += 1
            else:
                p.text = "và Tên môn: \t Mã MH: \t Số tín chỉ:"
        elif "Lý do:" in t:
            p.text = f"Lý do:  {st['reason']}"
            rec.append({"text": st['reason'], "class_name": "reason"})
        elif "ngày" in t and "tháng" in t and "năm" in t:
            p.text = f"                                                                      Tp. Hồ Chí Minh, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif "Kính đơn" in t:
            p.text = "                                                                                                        Kính đơn"
            p.paragraph_format.space_after = Pt(28)
        elif "Bộ môn quản môn học" in t:
            p.text = "Bộ môn quản môn học               Khoa quản môn học               Phòng Đào tạo"
        elif "(Ghi rõ ý kiến)" in t:
            p.text = "(Ghi rõ ý kiến: Đồng ý)"
            rec.append({"text": "Đồng ý", "class_name": "reason"})
            p.paragraph_format.space_after = Pt(24)
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)


def _fill_mien_thi(doc, st, rec):
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Họ tên sinh viên:" in t:
            p.text = (f"Họ tên sinh viên: {st['student_name']}    Điện thoại: {st['phone']}  "
                      f"Mã số sinh viên:  {st['student_id']} Lớp  {st['class_name']} Khoa  {st['faculty']}  "
                      f"Trường cũ đã học: {st['prev_university']}  "
                      f"Ngành cũ đã học: {st['prev_major']}")
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['phone'], "class_name": "phone"})
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
            rec.append({"text": st['prev_university'], "class_name": "degree_info"})
            rec.append({"text": st['prev_major'], "class_name": "major"})
        elif "Tp. HCM, ngày" in t:
            p.text = f"                                              Tp. HCM, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)

    if doc.tables:
        t = doc.tables[0]
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(8)
        courses = st["courses"][:min(len(st["courses"]), len(t.rows)-1)]
        for idx, c in enumerate(courses):
            row_idx = idx + 1
            if row_idx < len(t.rows):
                row = t.rows[row_idx]
                fill_table_cell(row.cells[0], str(idx+1), None, rec)
                fill_table_cell(row.cells[1], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[2], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[3], str(c["credits"]), "credits", rec, font_size_pt=8)
                fill_table_cell(row.cells[4], "Đạt yêu cầu", "reason", rec, font_size_pt=8)


def _fill_bao_luu(doc, st, rec):
    for sec in doc.sections:
        sec.top_margin = Inches(0.4)
        sec.bottom_margin = Inches(0.35)
        sec.left_margin = Inches(0.6)
        sec.right_margin = Inches(0.6)

    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Họ tên sinh viên:" in t and "Điện thoại:" in t:
            p.text = f"Họ tên sinh viên: {st['student_name']}        Điện thoại: {st['phone']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['phone'], "class_name": "phone"})
        elif "Mã số sinh viên:" in t and "Lớp" in t and "Khoa" in t:
            p.text = f"Mã số sinh viên: {st['student_id']}    Lớp: {st['class_name']}    Khoa: {st['faculty']}"
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Lớp cũ đã học:" in t and "Mã sinh viên cũ:" in t:
            p.text = f"Lớp cũ đã học: {st['prev_class']}    Mã sinh viên cũ: {st['prev_student_id']}"
            rec.append({"text": st['prev_class'], "class_name": "class_name"})
            rec.append({"text": st['prev_student_id'], "class_name": "student_id"})
        elif "Tp. HCM, ngày" in t:
            p.text = f"                                              Tp. HCM, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(8)
                        
        courses = st["courses"][:min(len(st["courses"]), 5)]
        for idx, c in enumerate(courses):
            row_idx = idx + 1
            if row_idx < len(t.rows):
                row = t.rows[row_idx]
                fill_table_cell(row.cells[0], str(idx+1), None, rec)
                fill_table_cell(row.cells[1], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[2], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[3], str(c["credits"]), "credits", rec, font_size_pt=8)
                fill_table_cell(row.cells[4], str(c["score"]), "score", rec, font_size_pt=8)
                fill_table_cell(row.cells[5], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[6], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[7], str(c["credits"]), "credits", rec, font_size_pt=8)
                fill_table_cell(row.cells[8], "Đồng ý", None, rec, font_size_pt=8)


def _fill_chuyen_diem_2018(doc, st, rec):
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Họ tên sinh viên:" in t and "Mã số sinh viên:" in t:
            p.text = f"Họ tên sinh viên: {st['student_name']}        Mã số sinh viên: {st['student_id']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['student_id'], "class_name": "student_id"})
        elif "Lớp:" in t and "Ngành:" in t and "Khoa:" in t:
            p.text = f"Lớp:  {st['class_name']}    Ngành:  {st['major']}    Khoa: {st['faculty']}    Điện thoại: {st['phone']}"
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['major'], "class_name": "major"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
            rec.append({"text": st['phone'], "class_name": "phone"})
        elif "Trường cũ đã học:" in t:
            p.text = f"Trường cũ đã học: {st['prev_university']}    Bậc: ĐH    Hệ: CQ"
            rec.append({"text": st['prev_university'], "class_name": "degree_info"})
        elif "Tp. Hồ Chí Minh, ngày" in t:
            p.text = f"                                              Tp. Hồ Chí Minh, ngày {st['day']} tháng {st['month']} Năm {st['year']}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} Năm {st['year']}", "class_name": "date"})
        elif "Xác nhận của Phòng Đào tạo" in t:
            p.text = "Xác nhận của Phòng Đào tạo                                         Sinh viên ký tên"
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)

    if doc.tables:
        t = doc.tables[0]
        # Format all rows in table to be compact
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(8)

        courses = st["courses"][:min(len(st["courses"]), 5)]
        for idx, c in enumerate(courses):
            row_idx = idx + 1
            if row_idx < len(t.rows):
                row = t.rows[row_idx]
                fill_table_cell(row.cells[0], f"{idx+1}.", None, rec)
                fill_table_cell(row.cells[1], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[2], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[3], str(c["credits"]), "credits", rec, font_size_pt=8)
                fill_table_cell(row.cells[4], str(c["score"]), "score", rec, font_size_pt=8)
                fill_table_cell(row.cells[5], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[6], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[7], str(c["credits"]), "credits", rec, font_size_pt=8)
                fill_table_cell(row.cells[8], "Đồng ý", None, rec, font_size_pt=8)


def _fill_chuyen_he(doc, st, rec):
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        t = p.text
        if "Đơn số:" in t:
            p.text = f"                                                                                                                     Đơn số: {random.randint(101, 999)}/ĐX-ĐT"
        elif "Em là:" in t and "Sinh ngày:" in t:
            p.text = f"Em là: {st['student_name']}        Sinh ngày: {st['dob']}        Tại: {st['pob']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['dob'], "class_name": "dob"})
            rec.append({"text": st['pob'], "class_name": "pob"})
        elif "MSSV:" in t and "Lớp:" in t and "Khoa:" in t:
            p.text = f"MSSV: {st['student_id']}    Lớp: {st['class_name']}    Khoa: {st['faculty']}"
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Điện thoại:" in t and "Email:" in t:
            p.text = f"Điện thoại: {st['phone']}    Email: {st['email']}"
            rec.append({"text": st['phone'], "class_name": "phone"})
            rec.append({"text": st['email'], "class_name": "email"})
        elif "Sinh viên ngành:" in t:
            p.text = f"Sinh viên ngành: {st['major']}"
            rec.append({"text": st['major'], "class_name": "major"})
        elif "Điểm tích lũy  toàn khóa:" in t or "Điểm tích lũy" in t:
            p.text = f"Điểm tích lũy  toàn khóa: {st['cpa']}    Tín chỉ tích lũy: {st['credits_accum']}"
            rec.append({"text": str(st['cpa']), "class_name": "score"})
            rec.append({"text": str(st['credits_accum']), "class_name": "credits"})
        elif "Bị buộc thôi học theo quyết định số" in t:
            p.text = f"Bị buộc thôi học theo quyết định số {st['book_number']} ngày {st['cert_date']}"
            rec.append({"text": st['book_number'], "class_name": "degree_info"})
            rec.append({"text": st['cert_date'], "class_name": "date"})
        elif "a/ Bậc cao đẳng" in t:
            p.text = f"a/ Bậc cao đẳng – ngành {st['second_major']} từ học kì 1 năm thứ 2"
            rec.append({"text": st['second_major'], "class_name": "major"})
        elif "Em xin cam đoan:" in t:
            p.text = "Em xin cam đoan:   -  Đã hoàn tất học phí và sách thư viện."
        elif "-………………………………………………………………………………" in t or ("-" in t and "……" in t):
            p.text = "     - Chấp hành nghiêm chỉnh mọi quy định của nhà trường."
        elif "Tp. Hồ Chí Minh, ngày" in t:
            p.text = f"                                                                    Tp. Hồ Chí Minh, ngày {st['day']}   tháng {st['month']} năm {st['year']}"
            rec.append({"text": f"ngày {st['day']}   tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif "Kính đơn" in t:
            p.text = "                                                                                   Kính đơn"
            p.paragraph_format.space_after = Pt(45)
        elif "Ban Giám hiệu duyệt" in t and "Đề xuất của P. Đào tạo" in t:
            p.text = " Ban Giám hiệu duyệt                                     Đề xuất của P. Đào tạo"
        elif "Đề nghị chuyển sinh viên" in t:
            p.text = f"Đề nghị chuyển sinh viên: {st['student_name']}"
            p.paragraph_format.left_indent = Inches(2.8)
            rec.append({"text": st['student_name'], "class_name": "student_name"})
        elif "Xuống bậc đào tạo" in t:
            p.text = f"Xuống bậc đào tạo: Cao đẳng   Ngành: {st['second_major']}"
            p.paragraph_format.left_indent = Inches(2.8)
            rec.append({"text": st['second_major'], "class_name": "major"})
        elif "HK…………...Năm thứ" in t or ("HK" in t and "Năm thứ" in t):
            p.text = "HK: 1                      Năm thứ: 2"
            p.paragraph_format.left_indent = Inches(2.8)
        elif "Tp. HCM, ngày…….tháng" in t or "Tp. HCM, ngày" in t:
            p.text = f"Tp. HCM, ngày {st['day']} tháng {st['month']} năm {st['year']}"
            p.paragraph_format.left_indent = Inches(2.8)
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm {st['year']}", "class_name": "date"})
        elif "Trưởng P.Đào tạo" in t:
            p.text = "Trưởng P.Đào tạo"
            p.paragraph_format.left_indent = Inches(3.4)
        elif not t.strip():
            p.text = ""
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(2)


def _fill_hoc_2_ctdt(doc, st, rec):
    for p in doc.paragraphs:
        t = p.text
        if "Tôi tên:" in t and "Sinh ngày:" in t:
            p.text = f"Tôi tên:   {st['student_name']}	Sinh ngày:  {st['dob']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['dob'], "class_name": "dob"})
        elif "Mã số sinh viên:" in t and "Lớp:" in t:
            p.text = f"Mã số sinh viên: {st['student_id']}	Lớp:  {st['class_name']}"
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['class_name'], "class_name": "class_name"})
        elif "Đang theo học ngành:" in t and "Khoa:" in t:
            p.text = (f"Đang theo học ngành: {st['major']}	Khoa:  {st['faculty']} "
                      f"Điểm trung bình chung học kỳ 1/năm học {st['academic_year']}:  {st['term_gpa']} "
                      f"Số tín chỉ đạt học kỳ: {st['credits_term']}")
            rec.append({"text": st['major'], "class_name": "major"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
            rec.append({"text": st['academic_year'], "class_name": "date"})
            rec.append({"text": str(st['term_gpa']), "class_name": "score"})
            rec.append({"text": str(st['credits_term']), "class_name": "credits"})
        elif "Điểm trung bình chung tích lũy:" in t:
            p.text = (f"Điểm trung bình chung tích lũy:  {st['cpa']}  "
                      f"Số tín chỉ tích lũy:  {st['credits_accum']} "
                      f"Tôi xin được học chương trình đào tạo thứ hai là: ")
            rec.append({"text": str(st['cpa']), "class_name": "score"})
            rec.append({"text": str(st['credits_accum']), "class_name": "credits"})
        elif "Tên CTĐT:" in t and "Khoa:" in t:
            p.text = f"Tên CTĐT:  {st['second_curriculum']}  Khoa:  {st['second_faculty']}"
            rec.append({"text": st['second_curriculum'], "class_name": "major"})
            rec.append({"text": st['second_faculty'], "class_name": "faculty"})
        elif "Tp.HCM, ngày" in t:
            p.text = f"	 	Tp.HCM, ngày {st['day']} tháng {st['month']} năm 20{st['year'][-2:]}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm 20{st['year'][-2:]}", "class_name": "date"})


def _fill_diem_i(doc, st, rec):
    for p in doc.paragraphs:
        t = p.text
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        if "Khoa (TT):" in t:
            p.text = f"			- Khoa (TT): {st['faculty']}"
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Họ và tên sinh viên :" in t or "Họ và tên sinh viên:" in t:
            p.text = f"Họ và tên sinh viên : {st['student_name']} Mã SV : {st['student_id']} Điện thoại: {st['phone']}"
            rec.append({"text": st['student_name'], "class_name": "student_name"})
            rec.append({"text": st['student_id'], "class_name": "student_id"})
            rec.append({"text": st['phone'], "class_name": "phone"})
        elif "Lớp :" in t or "Lớp:" in t:
            p.text = f"Lớp : {st['class_name']} Khoa : {st['faculty']}"
            rec.append({"text": st['class_name'], "class_name": "class_name"})
            rec.append({"text": st['faculty'], "class_name": "faculty"})
        elif "Xin được nhận điểm I" in t:
            p.text = f"Xin được nhận điểm I (điểm chưa hoàn tất) các môn học thuộc học kỳ 1 năm học {st['academic_year']} như sau:"
            rec.append({"text": st['academic_year'], "class_name": "date"})
        elif "Lý do :" in t or "Lý do:" in t:
            p.text = f"Lý do : {st['reason']}"
            rec.append({"text": st['reason'], "class_name": "reason"})
        elif "Tp.HCM, ngày" in t:
            p.text = f"Tp.HCM, ngày {st['day']} tháng {st['month']} năm 20{st['year'][-2:]}"
            rec.append({"text": f"ngày {st['day']} tháng {st['month']} năm 20{st['year'][-2:]}", "class_name": "date"})

    if len(doc.tables) > 1:
        t = doc.tables[1]
        num_courses = min(len(st["courses"]), random.choice([3, 4]))
        courses = st["courses"][:num_courses]
        for idx, c in enumerate(courses):
            row_idx = idx + 2
            if row_idx < len(t.rows):
                row = t.rows[row_idx]
                fill_table_cell(row.cells[0], str(idx+1), None, rec, font_size_pt=8)
                fill_table_cell(row.cells[1], c["code"], "course_code", rec, font_size_pt=7.5)
                fill_table_cell(row.cells[2], c["name"], "course_name", rec, font_size_pt=8)
                fill_table_cell(row.cells[3], c["group"], None, rec, font_size_pt=8)
                fill_table_cell(row.cells[4], c["exam_date"], "date", rec, font_size_pt=8)
                fill_table_cell(row.cells[5], c["room"], None, rec, font_size_pt=8)
                fill_table_cell(row.cells[6], c["period"], None, rec, font_size_pt=8)
                fill_table_cell(row.cells[7], "Đã ký", None, rec, font_size_pt=8)
                fill_table_cell(row.cells[8], str(c["score"]), "score", rec, font_size_pt=8)
                fill_table_cell(row.cells[9], "GV", None, rec, font_size_pt=8)


def _fill_generic(doc, st, rec):
    for p in doc.paragraphs:
        t = p.text
        if "..." in t or "…" in t:
            p.text = re.sub(r'[\.…_]{3,}', f" {st['student_name']} ", t, count=1)
            rec.append({"text": st['student_name'], "class_name": "student_name"})
