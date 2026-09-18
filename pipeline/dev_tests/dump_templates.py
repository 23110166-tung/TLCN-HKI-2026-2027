# -*- coding: utf-8 -*-
import docx, os, sys

sys.stdout.reconfigure(encoding='utf-8')
sample_dir = 'sample'
files = sorted([f for f in os.listdir(sample_dir) if f.endswith('.docx') and not f.startswith('~$')])

for f in files:
    print(f"\n============================== {f} ==============================")
    doc = docx.Document(os.path.join(sample_dir, f))
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            print(f"  P[{i:02d}]: {t}")
    for ti, t in enumerate(doc.tables):
        print(f"  TABLE {ti}: {len(t.rows)} rows x {len(t.columns)} cols")
        for ri, r in enumerate(t.rows[:2]):
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"    R[{ri}]: {cells}")
