# -*- coding: utf-8 -*-
import docx, os, sys

sys.stdout.reconfigure(encoding='utf-8')
sample_dir = 'sample'
files = sorted([f for f in os.listdir(sample_dir) if f.endswith('.docx') and not f.startswith('~$')])

for f in files:
    print(f"\n==================== {f} ====================")
    doc = docx.Document(os.path.join(sample_dir, f))
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"  P[{i}]: {p.text}")
            for ri, r in enumerate(p.runs):
                if r.text.strip():
                    print(f"    R[{ri}]: '{r.text}'")
    for ti, t in enumerate(doc.tables):
        print(f"  TABLE {ti} ({len(t.rows)}x{len(t.columns)}):")
        for ri, r in enumerate(t.rows):
            cells_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"    Row {ri}: {cells_txt}")
